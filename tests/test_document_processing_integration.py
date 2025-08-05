import asyncio
import logging
import sys
from pathlib import Path

import pytest

# Add paths for the various pipelines
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "AI_SCRAPER_RAW_TEXTS"))
sys.path.insert(0, str(ROOT / "AgriTool-Document-Schema-Extractor"))

from document_processing import PythonDocumentExtractor, ScraperDocumentExtractor
from AI_SCRAPER_RAW_TEXTS.scraper.extract_raw_page import RawPageExtractor
from document_schema_extractor import SharedDocumentExtractor


def _create_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello DOCX")
    doc.save(path)
    return path


def _create_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello XLSX"
    wb.save(path)
    return path


def _create_pdf(path: Path) -> Path:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.drawString(100, 750, "Hello PDF")
    c.save()
    return path


@pytest.mark.parametrize(
    "factory",
    [
        lambda: PythonDocumentExtractor(),
        lambda: ScraperDocumentExtractor().extractor,
        lambda: RawPageExtractor(driver=None, output_dir="test_output").doc_extractor,
    ],
)
def test_pdf_docx_xlsx_extraction(tmp_path, factory):
    pdf_path = _create_pdf(tmp_path / "sample.pdf")
    docx_path = _create_docx(tmp_path / "sample.docx")
    xlsx_path = _create_xlsx(tmp_path / "sample.xlsx")

    extractor = factory()

    for path in [pdf_path, docx_path, xlsx_path]:
        result = extractor.extract_document(str(path))
        assert result.success
        assert result.text_content.strip()


def test_schema_extractor_uses_shared_module(tmp_path):
    pdf_path = _create_pdf(tmp_path / "sample.pdf")
    docx_path = _create_docx(tmp_path / "sample.docx")
    xlsx_path = _create_xlsx(tmp_path / "sample.xlsx")

    extractor = SharedDocumentExtractor(logger=logging.getLogger("test"))

    async def run():
        for path in [pdf_path, docx_path, xlsx_path]:
            data = await extractor.extract_schema(str(path))
            assert data["raw_unclassified"][0].strip()

    asyncio.run(run())
